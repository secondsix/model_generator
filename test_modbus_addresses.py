import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from modbus_addresses import convert_modbus_address, is_modbus_file
from point_generator import POINT_COLUMNS, generate_point_file, extract_address


class ModbusTests(unittest.TestCase):
    def test_conversion_and_reference_overrides(self):
        filename = "ZS00990(ZS330 D21 Modbus).xlsx"
        self.assertTrue(is_modbus_file(filename))
        self.assertTrue(is_modbus_file("device.MODBUS.xlsx"))
        self.assertFalse(is_modbus_file("device.xlsx"))
        for raw, comment, expected in (("3x400525", "设定射胶1段压力", "HR524"), ("3x401512", "设定保压1段时间", "HR1512"), ("3x401007", "实际射嘴温度", "HR1007"), ("3x401573", "实际冷却时间", "HR1572"), ("0x1.0", "电机", "CO0"), ("0x16", "报警灯", "CO15")):
            self.assertEqual(convert_modbus_address(raw, comment, filename)[0], expected)
        self.assertTrue(convert_modbus_address("3x401512", "设定保压1段时间", filename)[1])
        self.assertEqual(convert_modbus_address("3x401512", "设定保压1段时间", "other Modbus.xlsx")[0], "HR1511")
        self.assertEqual(convert_modbus_address("3x401545", "实际熔胶时间", filename)[0], "")
        for raw in ("bad", "0x0", "3x400000", "3x499999"):
            self.assertEqual(convert_modbus_address(raw, "unknown", filename)[0], "")
        self.assertEqual(extract_address('"ns=2;i=20250","6;-1",0'), "ns=2;i=20250")

    def test_modbus_export(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, device, template = root / "单色工艺参数.xlsx", root / "ZS00990 D21 Modbus.xlsx", root / "template.xlsx"
            fixtures = [(source, [["工艺参数名称", "工艺参数单位"], ["保压时间一段", "s"], ["注射压力一段", "bar"]]), (device, [["注释", "地址", "数据类型"], ["设定保压1段时间", "3x401512", "无符号的 32 位值"], ["设定射胶1段压力", "3x400525", "无符号的 16 位值"]]), (template, [list(POINT_COLUMNS)])]
            for path, rows in fixtures:
                wb = Workbook()
                for row in rows: wb.active.append(row)
                wb.save(path)
                wb.close()
            result = generate_point_file(source, device, template, root / "output")
            self.assertEqual(result["matched_count"], 2)
            self.assertIn("Modbus", result["rule"])
            wb = load_workbook(result["generated_files"][0])
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual(rows[0][2:4], ("HR1512", "UINT32"))
            self.assertEqual(rows[1][2:4], ("HR524", "UINT16"))
            self.assertEqual(rows[0][5:], (None, 1000, None, None, None))

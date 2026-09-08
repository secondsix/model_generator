import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from tag_generator import TAG_COLUMNS, generate_tag_files, make_initials
from model_generator import SOURCE_COLUMNS, read_models
from model_generator_gui import ModelGeneratorApp


class TagGeneratorTests(unittest.TestCase):
    def test_duplicate_names_across_products(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, template = root / "source.xlsx", root / "template.xlsx"
            wb = Workbook()
            wb.active.append(list(TAG_COLUMNS))
            wb.save(template)
            wb.close()
            entries = [("移印", "温度", "℃"), ("硬压", "温度", "℃"), ("硬压", "压力", "bar")]
            for order in (entries, entries[::-1]):
                wb = Workbook()
                wb.active.append(list(SOURCE_COLUMNS.values()))
                for product, name, unit in order:
                    wb.active.append([product, "运行参数", name, unit, "说明"])
                wb.save(source)
                wb.close()
                result = generate_tag_files(source, template, root / "output")
                actual = {}
                for path in result["generated_files"]:
                    wb = load_workbook(path)
                    actual.update({row[2]: row[3:] for row in list(wb.active.values)[1:]})
                    wb.close()
                self.assertEqual(actual, {
                    "GY_YIYIN_WD": ("温度_移印", "说明", "℃"),
                    "GY_YINGYA_WD": ("温度_硬压", "说明", "℃"),
                    "GY_YINGYA_YL": ("压力", "说明", "bar"),
                })

    def test_full_categories_and_separate_product_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, template = root / "source.xlsx", root / "template.xlsx"
            wb = Workbook()
            wb.active.append(list(SOURCE_COLUMNS.values()))
            wb.active.append(["移印", "运行参数", "温度", None, None])
            wb.active.append(["硬压", "运行参数", "压力", None, None])
            wb.active.append(["激光打码", "运行参数", "频率", None, None])
            wb.save(source)
            wb.close()
            wb = Workbook()
            wb.active.append(list(TAG_COLUMNS))
            wb.save(template)
            wb.close()
            progress = []
            result = generate_tag_files(source, template, root / "output", progress=lambda *args: progress.append(args))
            self.assertEqual(len(result["generated_files"]), 3)
            self.assertEqual(result["tag_count"], 3)
            for product, category, name, suffix in (("移印", "YIYIN", "温度", "WD"), ("硬压", "YINGYA", "压力", "YL"), ("激光打码", "JGDM", "频率", "PL")):
                output = root / "output" / f"{product}物模型标签库导入.xlsx"
                self.assertIn(output, result["generated_files"])
                wb = load_workbook(output)
                rows = list(wb.active.values)
                wb.close()
                self.assertEqual(len(rows), 2)
                self.assertEqual(rows[1], ("property", category, f"GY_{category}_{suffix}", name, None, None))
            self.assertEqual([event[:2] for event in progress], [(1, 3), (2, 3), (3, 3)])

    def test_initials(self):
        self.assertEqual(make_initials("注射压力额定"), "ZSYLED")
        self.assertEqual(make_initials("LED温度2"), "LEDWD2")

    def test_export_and_collision(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, template, output = (root / name for name in ("source.xlsx", "template.xlsx", "nested/双色注塑机物模型标签库导入.xlsx"))
            wb = Workbook()
            ws = wb.active
            ws.append(list(SOURCE_COLUMNS.values()))
            ws.append(["双色注塑机", "运行参数", "注射压力额定", "MPa", "压力说明"])
            ws.append([None, "运行参数", "A/B组温度", None, None])
            ws.append([None, "服务", "不导出", None, None])
            wb.create_sheet("忽略").append(["其他内容"])
            wb.save(source)
            wb.close()
            for with_example in (False, True):
                wb = Workbook()
                wb.active.append(list(TAG_COLUMNS))
                if with_example:
                    wb.active.append(["property", "OLD", "GY_OLD_OLD", "旧示例", "旧说明"])
                    wb.active.append(["property", "OLD", "GY_OLD_OLD2", "旧示例2", None])
                    wb.active.append(["property", "OLD", "GY_OLD_OLD3", "旧示例3", None])
                wb.save(template)
                wb.close()
                result = generate_tag_files(source, template, output.parent)
                self.assertEqual(result["tag_count"], 2)
                wb = load_workbook(output)
                rows = list(wb.active.values)
                wb.close()
                self.assertEqual(rows[1], ("property", "SSZSJ", "GY_SSZSJ_ZSYLED", "注射压力额定", "压力说明", "MPa"))
                self.assertEqual(rows[2][3:], ("A/B组温度", None, None))
                self.assertEqual(len(rows), 3)
            self.assertEqual(len(read_models(source)[1]["双色注塑机"]["parameters"]), 3)
            original = output.read_bytes()
            wb = load_workbook(source)
            wb.active.append([None, "运行参数", "注射压力额定", None, None])
            wb.save(source)
            wb.close()
            with self.assertRaisesRegex(ValueError, "标识冲突"):
                generate_tag_files(source, template, output.parent)
            self.assertEqual(output.read_bytes(), original)
            with self.assertRaisesRegex(ValueError, "文件夹"):
                generate_tag_files(source, template, source)

    def test_ui_mode_and_save_dialog(self):
        app = ModelGeneratorApp()
        app.withdraw()
        try:
            app.mode_var.set("物模型导入文件生成器")
            app._switch_mode()
            selected = ("C:/my files/甲.xlsx", "C:/my files/乙.xlsx")
            with patch("model_generator_gui.filedialog.askopenfilenames", return_value=selected):
                app._choose_source()
            self.assertEqual(app._source_paths(), list(map(Path, selected)))
            app.template_var.set("old-template.xlsx")
            app.output_var.set("old-folder")
            app.mode_var.set("物模型标签库导入生成器")
            app._switch_mode()
            with patch("model_generator_gui.filedialog.askdirectory", return_value="C:/chosen") as dialog:
                app._choose_output()
                dialog.assert_called_once()
            self.assertEqual(app.output_var.get(), "C:/chosen")
            app.mode_var.set("物模型导入文件生成器")
            app._switch_mode()
            self.assertEqual(app.output_var.get(), "old-folder")
            self.assertEqual(app.template_var.get(), "old-template.xlsx")
            self.assertEqual(app._source_paths(), list(map(Path, selected)))
        finally:
            app.destroy()

    def test_number_collision_export_is_order_independent(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, template, output = (root / name for name in ("source.xlsx", "template.xlsx", "单色注塑物模型标签库导入.xlsx"))
            wb = Workbook()
            wb.active.append(list(TAG_COLUMNS))
            wb.save(template)
            wb.close()
            names = ["注射压力三段", "注射压力四段", "注射压力十段", "注射压力二段"]
            expected = {
                names[0]: "GY_DSZS_ZSYLSAND", names[1]: "GY_DSZS_ZSYLSID",
                names[2]: "GY_DSZS_ZSYLSHID", names[3]: "GY_DSZS_ZSYLED",
            }
            for order in (names, names[::-1]):
                wb = Workbook()
                wb.active.append(list(SOURCE_COLUMNS.values()))
                for name in order:
                    wb.active.append(["单色注塑", "运行参数", name, None, None])
                wb.save(source)
                wb.close()
                generate_tag_files(source, template, output.parent)
                wb = load_workbook(output)
                actual = {row[3]: row[2] for row in list(wb.active.values)[1:]}
                wb.close()
                self.assertEqual(actual, expected)

    def test_duplicate_names_use_units(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, template, output = (root / name for name in ("source.xlsx", "template.xlsx", "激光打码物模型标签库导入.xlsx"))
            wb = Workbook()
            wb.active.append(list(TAG_COLUMNS))
            wb.save(template)
            wb.close()
            entries = [("打点时间", "μs"), ("打点时间", "ms"), ("频率", "Hz"), ("频率", "KHz"), ("延时", "ms")]
            expected = {
                "打点时间微秒": "GY_JGDM_DDSJWM", "打点时间毫秒": "GY_JGDM_DDSJHM",
                "频率赫兹": "GY_JGDM_PLHZ", "频率千赫兹": "GY_JGDM_PLQHZ",
                "延时": "GY_JGDM_YS",
            }
            for order in (entries, entries[::-1]):
                wb = Workbook()
                wb.active.append(list(SOURCE_COLUMNS.values()))
                for name, unit in order:
                    wb.active.append(["激光打码", "运行参数", name, unit, "原始说明"])
                wb.save(source)
                wb.close()
                generate_tag_files(source, template, output.parent)
                wb = load_workbook(output)
                rows = list(wb.active.values)[1:]
                wb.close()
                self.assertEqual({row[3]: row[2] for row in rows}, expected)
                self.assertTrue(all(row[4] == "原始说明" for row in rows))
            original = output.read_bytes()
            wb = load_workbook(source)
            wb.active.append(["激光打码", "运行参数", "打点时间", "µs", None])
            wb.save(source)
            wb.close()
            with self.assertRaisesRegex(ValueError, "标识冲突"):
                generate_tag_files(source, template, output.parent)
            self.assertEqual(output.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()

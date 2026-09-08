"""按标签库模板为每个产品生成独立导入文件。"""

import os
import re
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from pypinyin import Style, lazy_pinyin

from model_generator import copy_row_style, find_header_row, get_header_map, read_models, safe_filename

TAG_COLUMNS = ("标签类型", "分类标识", "标识", "名称", "说明", "单位")
NUMBER_PINYIN = dict(zip("零〇一二两三四五六七八九十百千万亿", (
    "LING", "LING", "YI", "ER", "LIANG", "SAN", "SI", "WU", "LIU", "QI",
    "BA", "JIU", "SHI", "BAI", "QIAN", "WAN", "YI")))
UNIT_NAMES = {
    "μs": "微秒", "µs": "微秒", "us": "微秒", "ms": "毫秒",
    "s": "秒", "ns": "纳秒", "Hz": "赫兹", "hz": "赫兹",
    "kHz": "千赫兹", "KHz": "千赫兹", "khz": "千赫兹",
    "MHz": "兆赫兹", "GHz": "吉赫兹",
}


def make_initials(name):
    """中文取拼音首字母，英文和数字保留，统一大写。"""
    result = re.sub(r"[^A-Za-z0-9]", "", "".join(lazy_pinyin(name, style=Style.FIRST_LETTER))).upper()
    if not result:
        raise ValueError(f"无法从名称生成标识：{name}")
    return result


def expand_number_initials(name):
    """冲突时将中文数字展开为全拼，其余部分仍取首字母。"""
    return make_initials("".join(NUMBER_PINYIN.get(char, char) for char in name))


def make_category(name):
    result = re.sub(r"[^A-Za-z0-9]", "", "".join(lazy_pinyin(name))).upper()
    if not result:
        raise ValueError(f"无法从产品名称生成分类标识：{name}")
    return result


def generate_tag_files(source_file, template_file, output_dir, progress=None):
    source_file, template_file, output_dir = map(lambda p: Path(p).resolve(), (source_file, template_file, output_dir))
    if output_dir.exists() and not output_dir.is_dir():
        raise ValueError("请选择文件夹作为导出位置。")
    sheet_name, models, included, skipped = read_models(source_file, expand_names=False)
    planned, categories, identifiers, filenames = [], {}, {}, set()
    category_counts = Counter(make_initials(product) for product in models)
    for product, model in models.items():
        category = make_initials(product)
        if category_counts[category] > 1:
            category = make_category(product)
        rows = []
        if category in categories and categories[category] != product:
            raise ValueError(f"分类标识冲突 {category}：{categories[category]}、{product}")
        categories[category] = product
        parameters = model["parameters"]
        name_counts = Counter(p["name"] for p in parameters)
        missing_units = {p["name"] for p in parameters if not p["unit"].strip()}
        export_names = []
        for parameter in parameters:
            name = parameter["name"]
            if name_counts[name] > 1 and name not in missing_units:
                unit = parameter["unit"].strip()
                name += UNIT_NAMES.get(unit, unit)
            export_names.append(name)
        initial_counts = Counter(make_initials(name) for name in export_names)
        for parameter, name in zip(parameters, export_names):
            initials = make_initials(name)
            if initial_counts[initials] > 1:
                initials = expand_number_initials(name)
            identifier = f"GY_{category}_{initials}"
            if identifier in identifiers:
                raise ValueError(f"标签标识冲突 {identifier}：{identifiers[identifier]}、{product}/{name}")
            identifiers[identifier] = f"{product}/{name}"
            rows.append(("property", category, identifier, name, parameter["description"], parameter["unit"]))
        output_file = output_dir / f"{safe_filename(product)}物模型标签库导入.xlsx"
        if output_file in (source_file, template_file):
            raise ValueError("保存位置不能与设计表或模板相同。")
        key = output_file.name.casefold()
        if key in filenames:
            raise ValueError(f"产品名称生成的文件名重复：{output_file.name}")
        filenames.add(key)
        planned.append((product, output_file, rows))
    if not planned:
        raise ValueError("设计表第一个 sheet 中没有可导出的参数记录。")
    # 在全部产品间检查最终名称（已包含单位后缀），只修改名称列。
    export_name_counts = Counter(row[3] for _, _, rows in planned for row in rows)
    used_names = set()
    for product, _, rows in planned:
        for index, row in enumerate(rows):
            name = row[3]
            if export_name_counts[name] > 1:
                name = f"{name}_{product}"
            if name in used_names:
                raise ValueError(f"追加产品名称后标签名称仍重复：{name}")
            used_names.add(name)
            rows[index] = (*row[:3], name, *row[4:])
    for index, (product, output_file, rows) in enumerate(planned, 1):
        write_tag_file(template_file, output_file, rows)
        if progress:
            progress(index, len(planned), product, output_file)
    return {"sheet_name": sheet_name, "included_types": included, "skipped_types": skipped,
            "generated_files": [item[1] for item in planned], "output_dir": output_dir,
            "tag_count": sum(len(item[2]) for item in planned)}


def write_tag_file(template_file, output_file, rows):
    wb = load_workbook(template_file)
    temporary = None
    try:
        ws = wb.worksheets[0]
        header = find_header_row(ws, TAG_COLUMNS[:-1])
        headers = get_header_map(ws, header)
        if "单位" not in headers:
            column = ws.max_column + 1
            ws.cell(header, column, "单位")
            ws.column_dimensions[ws.cell(header, column).column_letter].width = 14
            headers["单位"] = column
        for row in ws.iter_rows(min_row=header + 1):
            for cell in row:
                cell.value = None
        for index, values in enumerate(rows, header + 1):
            if index != header + 1:
                copy_row_style(ws, header + 1, index)
            for field, value in zip(TAG_COLUMNS, values):
                cell = ws.cell(index, headers[field], value)
                cell.data_type = "s"
        last = header + len(rows)
        if ws.max_row > last:
            ws.delete_rows(last + 1, ws.max_row - last)
        output_file.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=output_file.parent, suffix=".xlsx", delete=False) as temp:
            temporary = Path(temp.name)
        wb.save(temporary)
        os.replace(temporary, output_file)
    finally:
        wb.close()
        if temporary and temporary.exists():
            temporary.unlink()

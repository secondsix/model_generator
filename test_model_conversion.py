import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from model_generator import ALLOWED_DATA_TYPES, TEMPLATE_COLUMNS, SOURCE_COLUMNS, generate_files, generate_many_files, get_data_type
from tag_generator import generate_tag_files


class ConversionTests(unittest.TestCase):
    def test_multiple_libraries(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "template.xlsx"
            wb = Workbook()
            wb.active.append(list(TEMPLATE_COLUMNS))
            wb.save(template)
            wb.close()
            sources = []
            for name in ("产品 A", "产品 B"):
                path = root / f"{name}物模型标签库导入.xlsx"
                wb = Workbook()
                wb.active.append(["标识", "名称", "单位"])
                wb.active.append([name, "压力", "bar"])
                wb.save(path)
                wb.close()
                sources.append(path)
            events = []
            result = generate_many_files(sources, template, root / "output", progress=lambda *args: events.append(args))
            self.assertEqual(len(result["generated_files"]), 2)
            self.assertEqual([event[:2] for event in events], [(1, 2), (2, 2)])
            for path, name in zip(result["generated_files"], ("产品 A", "产品 B")):
                wb = load_workbook(path)
                self.assertEqual(wb.active.cell(2, 1).value, name)
                wb.close()
            duplicate = root / "产品 A.xlsx"
            duplicate.write_bytes(sources[0].read_bytes())
            with self.assertRaisesRegex(ValueError, "文件名重复"):
                generate_many_files([sources[0], duplicate], template, root / "conflict")
            self.assertFalse((root / "conflict").exists())

    def test_unit_rules(self):
        for unit in ("bar", "mm/s", "%", "s", "r/min", "mm", "℃", "μs", "ms", "Hz", "KHz"):
            info = get_data_type("状态时间戳", unit)
            self.assertEqual(info["data_type"], "double")
            self.assertEqual(info["unit"], unit)
            self.assertEqual(json.loads(info["config"])["scale"], info["precision"])
        self.assertEqual(get_data_type("数量", "个")["data_type"], "int")
        self.assertEqual(get_data_type("累计数量", "个")["data_type"], "long")
        self.assertEqual(get_data_type("编号", "")["data_type"], "string")
        self.assertEqual(get_data_type("是否故障", "")["data_type"], "boolean")
        for kind in ALLOWED_DATA_TYPES:
            info = get_data_type("测试", kind)
            self.assertEqual(info["data_type"], kind)
            self.assertEqual(json.loads(info["config"])["type"], kind)

    def test_design_to_tags_to_model(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, tags_template, model_template = (root / name for name in ("design.xlsx", "tags.xlsx", "model.xlsx"))
            for path, rows in (
                (source, [list(SOURCE_COLUMNS.values()), ["激光打码", "运行参数", "打点时间", "μs", "说明一"], [None, "运行参数", "打点时间", "ms", "说明二"]]),
                (tags_template, [["标签类型", "分类标识", "标识", "名称", "说明"]]),
                (model_template, [list(TEMPLATE_COLUMNS), ["old", "示例", "string"]]),
            ):
                wb = Workbook()
                for row in rows:
                    wb.active.append(row)
                wb.save(path)
                wb.close()
            tag_path = generate_tag_files(source, tags_template, root / "tags")["generated_files"][0]
            wb = load_workbook(tag_path)
            tags = list(wb.active.values)
            wb.close()
            self.assertEqual(tags[0][-1], "单位")
            self.assertEqual([row[-1] for row in tags[1:]], ["μs", "ms"])
            model_path = generate_files(tag_path, model_template, root / "models")["generated_files"][0]
            wb = load_workbook(model_path)
            rows = list(wb.active.values)
            wb.close()
            self.assertEqual(len(rows), 3)
            for tag, row in zip(tags[1:], rows[1:]):
                self.assertEqual(row[:2], (tag[2], tag[3]))
                self.assertEqual(row[2:5], ("double", tag[5], 2))
                self.assertEqual(row[6:], ("设备", tag[4], "上报", "存储"))
                self.assertEqual(json.loads(row[5])["unit"], row[3])


if __name__ == "__main__":
    unittest.main()

"""注塑机工艺参数与 WinCC 点表匹配生成器。"""
import csv
import os
import re
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook
from model_generator import clean_text, find_header_row, get_header_map, copy_row_style
from tag_generator import make_initials, expand_number_initials
from modbus_addresses import is_modbus_file, convert_modbus_address

POINT_COLUMNS = ("点位名称", "点位标识", "协议地址", "数据类型", "单位", "点位标签", "采集周期(ms)", "协议配置(JSON)", "输出类型", "转换配置(JSON)")
POINT_TYPES = {"BOOL", "INT8", "INT16", "INT32", "INT64", "UINT8", "UINT16", "UINT32", "UINT64", "FLOAT32", "FLOAT64", "STRING", "BYTES", "DATETIME"}
NUMBERS = "一二三四五六七八九十"


def normalize(text):
    return re.sub(r"\s+|<br\s*/?>", "", clean_text(text), flags=re.I).replace("（", "(").replace("）", ")")


def single_fallback(name):
    """第二套单色规则，仅作为已有候选之后的备用注释。"""
    normalized = normalize(name)
    base = re.sub(r"\([^)]*\)", "", normalized)
    rules = {"注射时间": "射出时间", "冷却时间": "冷却时间",
             "切换时间": "切保压时间(射胶时间)", "切换保压位置": "切保压位置",
             "注射终点位置": "射出终点位置",
             "储料速度(mm/s)": "储料速度1", "储料速度(r/min)": "储料速度2",
             "储料速度(%)": "储料速度3"}
    for stem, count in (("注射压力", 4), ("保压压力", 4), ("注射速度", 6), ("保压时间", 4), ("储料位置", 3)):
        for index, number in enumerate(NUMBERS[:count], 1):
            rules[f"{stem}{number}段"] = f"{stem}{index}"
    return rules.get(normalized, rules.get(base, ""))


def match_names(name, dual=False):
    base = re.sub(r"\([^)]*\)", "", normalize(name))
    group = ""
    if dual:
        match = re.match(r"([AB])组(.*)", base)
        if not match:
            return [base]
        group, base = match.groups()
    rules = {
        "背压": "储料背压", "注射时间": "注射+保压时间", "螺杆转速": "螺杆转速",
        "保压速度": "保压速度", "冷却时间": "冷却实际时间" if dual else "冷却",
        "切换时间": "V/P切换时间", "切换保压位置": "V/P切换位置",
        "储料速度": {"mm/s": "储料1段速度", "r/min": "储料2段速度", "%": "储料3段速度"}.get(
            re.search(r"\(([^)]*)\)", normalize(name)).group(1) if "(" in normalize(name) else "", "储料速度"),
        "射退距离": "射退", "加料时间": "储料时间", "周期": "循环周期时间",
        "注射终点位置": "注射最前位置", "射出终点位置": "注射最前位置",
    }
    target = rules.get(base)
    for index, number in enumerate(NUMBERS, 1):
        stage = number if dual else str(index)
        staged = {f"注射压力{number}段": f"射出{stage}段压力",
                  f"注射速度{number}段": f"射出{stage}段速度",
                  f"保压压力{number}段": f"保压第{number}段压力" if dual else f"保压第段压力{index}",
                  f"保压时间{number}段": f"保压第{number}段时间" if dual else f"保压第段时间{index}",
                  f"储料位置{number}段": f"储料{stage}段位置",
                  f"料筒温度{number}段": f"{number}段温度实际",
                  f"机筒温度{number}段": f"{number}段温度实际"}
        if base in staged:
            target = staged[base]
            break
    # 原有候选保留组别；备用规则仅对用户指定项目使用无组别注释。
    candidates = [target + group] if target else []
    candidates += [base + group, normalize(name)]
    if not dual:
        fallback = single_fallback(name)
        if fallback:
            candidates.append(fallback)
        third_rules = {"冷却时间": "冷却时间", "切换时间": "注射切换时间", "余料位置": "余料位置"}
        for index, number in enumerate(NUMBERS[:3], 1):
            for unit in ("mm/s", "%"):
                third_rules[f"注射速度{number}段({unit})"] = f"注射速度{index}段"
        third = third_rules.get(normalize(name))
        if third:
            candidates.append(third)
        fourth_rules = {"注射时间": "射出时间", "切换保压位置": "转保压位置",
                        "注射终点位置": "射出终点", "射退距离": "射退距离"}
        for index, number in enumerate(NUMBERS[:6], 1):
            fourth_rules[f"注射速度{number}段(mm/s)"] = f"射出{index}段速度"
            fourth_rules[f"注射速度{number}段(%)"] = f"射出{index}段速度B"
            if index <= 4:
                fourth_rules[f"注射压力{number}段"] = f"射出{index}段压力"
                fourth_rules[f"保压压力{number}段"] = f"保压{index}段压力"
            if index <= 3:
                fourth_rules[f"储料位置{number}段"] = f"储料{index}段位置"
        fourth = fourth_rules.get(normalize(name))
        if fourth:
            candidates.append(fourth)
        fifth_rules = {
            "背压": "ZS450B_背压设定值", "注射时间": "ZS450B_注射时间",
            "螺杆转速(r/min)": "ZS450B_螺杆转速设定", "螺杆转速(%)": "ZS450B_螺杆转速设定",
            "冷却时间": "ZS450B_冷却时间设定值", "加料时间": "ZS450B_加料时间实际值",
            "切换时间": "实际值PN切换时间", "切换保压位置": "ZS450B_切换位置实际值",
            "余料位置": "ZS450B_余料位置", "射嘴温度": "ZS450B_射嘴实际温度",
            "喷嘴温度": "实际值温度喷嘴1",
        }
        for index, number in enumerate(NUMBERS[:6]):
            for unit in ("mm/s", "%"):
                fifth_rules[f"注射速度{number}段({unit})"] = "设定值注入速度" + (str(index) if index else "")
        fifth = fifth_rules.get(normalize(name))
        if fifth:
            candidates.append(fifth)
        sixth_rules = {
            "注射压力一段": "ZS474双色_注射压力",
            "背压": "ZS474双色_背压设定值",
            "注射速度一段(mm/s)": "ZS474双色_注射速度",
            "注射时间": "ZS474双色_注射时间",
            "螺杆转速(r/min)": "ZS474双色_螺杆转速设定",
            "螺杆转速(%)": "ZS474双色_螺杆转速设定",
            "冷却时间": "ZS474双色_冷却时间设定值",
            "加料时间": "ZS474双色_加料时间",
            "切换保压位置": "ZS474双色_切换位置实际值",
            "余料位置": "ZS474双色_余料位置",
            "射嘴温度": "ZS474SS_射嘴实际温度",
        }
        sixth = sixth_rules.get(normalize(name))
        if sixth:
            candidates.append(sixth)
        seventh_rules = {
            "注射压力一段": "ZS451双色_注射压力",
            "保压压力一段": "ZS451双色_保压压力",
            "背压": "ZS451双色_背压设定值",
            "注射速度一段(mm/s)": "ZS451双色_注射速度",
            "注射速度一段(%)": "ZS451双色_注射速度",
            "注射时间": "ZS451双色_注射时间",
            "螺杆转速(r/min)": "ZS451双色_螺杆转速设定",
            "螺杆转速(%)": "ZS451双色_螺杆转速设定",
            "保压时间一段": "ZS451双色_保压时间",
            "冷却时间": "ZS451双色_冷却时间设定值",
            "加料时间": "ZS451双色_加料时间",
            "切换保压位置": "ZS451双色_切换位置",
            "余料位置": "ZS451双色_余料位置",
            "射嘴温度": "ZS451双色_射嘴实际温度",
        }
        seventh = seventh_rules.get(normalize(name))
        if seventh:
            candidates.append(seventh)
        eighth_rules = {"冷却时间": "实际冷却时间", "射嘴温度": "实际射嘴温度"}
        for index, number in enumerate(NUMBERS[:4], 1):
            eighth_rules[f"注射压力{number}段"] = f"设定射胶{index}段压力"
            eighth_rules[f"保压压力{number}段"] = f"设定保压{index}段压力"
            eighth_rules[f"保压时间{number}段"] = f"设定保压{index}段时间"
        for index, number in enumerate(NUMBERS[:5], 1):
            for unit in ("mm/s", "%"):
                eighth_rules[f"注射速度{number}段({unit})"] = f"设定射胶{index}段速度"
        eighth = eighth_rules.get(normalize(name))
        if eighth:
            candidates.append(eighth)
    else:
        dual_rules = {"冷却时间": f"{group}组冷却时间",
                      "加料时间": "储料时间" if group == "A" else "B储料时间",
                      "切换时间": f"{group}组切保压时间(射胶时间)",
                      "切换保压位置": f"{group}组切保压位置",
                      "射出终点位置": "射出终点位置"}
        for stem, count in (("注射压力", 4), ("保压压力", 4), ("注射速度", 6), ("保压时间", 4), ("储料位置", 3)):
            for index, number in enumerate(NUMBERS[:count], 1):
                dual_rules[f"{stem}{number}段"] = f"{group}组{stem}{index}"
        fallback = dual_rules.get(base)
        if fallback:
            candidates.append(fallback)
    return list(dict.fromkeys(candidates))


def convert_type(value):
    value = normalize(value)
    upper = value.upper()
    if upper in POINT_TYPES:
        return upper
    if "无符号" in value or "有符号" in value:
        bits = re.search(r"(8|16|32|64)位", value)
        if bits:
            return ("UINT" if "无符号" in value else "INT") + bits.group(1)
    if "浮点" in value:
        if "64" in value: return "FLOAT64"
        if "32" in value: return "FLOAT32"
    return {"BOOL": "BOOL", "BOOLEAN": "BOOL", "二进制变量": "BOOL", "布尔": "BOOL",
            "REAL": "FLOAT32", "LREAL": "FLOAT64", "文本变量8位字符集": "STRING",
            "文本变量16位字符集": "STRING", "字符串": "STRING", "字节数组": "BYTES",
            "日期时间": "DATETIME", "DATE_AND_TIME": "DATETIME"}.get(value, "")


def extract_address(value):
    text = clean_text(value)
    if not text:
        return ""
    first = next(csv.reader([text]))[0].strip()
    return first if re.fullmatch(r"ns=\d+;[isgb]=.+", first) else ""


def generate_point_file(parameter_file, device_file, template_file, output_dir, progress=None):
    parameter_file, device_file, template_file, output_dir = map(lambda p: Path(p).resolve(), (parameter_file, device_file, template_file, output_dir))
    modbus = is_modbus_file(device_file)
    if "双色" in parameter_file.stem:
        dual = True
    elif "单色" in parameter_file.stem:
        dual = False
    else:
        raise ValueError("工艺参数文件名需包含“单色”或“双色”，以确定匹配规则。")
    wb = load_workbook(parameter_file, data_only=True)
    try:
        ws = wb.worksheets[0]
        header = find_header_row(ws, ("工艺参数名称", "工艺参数单位"))
        columns = get_header_map(ws, header)
        parameters = [(clean_text(ws.cell(r, columns["工艺参数名称"]).value), clean_text(ws.cell(r, columns["工艺参数单位"]).value)) for r in range(header + 1, ws.max_row + 1)]
        parameters = [(name, unit) for name, unit in parameters if name]
    finally:
        wb.close()
    if not parameters:
        raise ValueError("没有可导出的工艺参数。")
    wb = load_workbook(device_file, data_only=True)
    lookup = defaultdict(set)
    try:
        ws = wb.worksheets[0]
        header = find_header_row(ws, ("注释", "地址", "数据类型"))
        columns = get_header_map(ws, header)
        for r in range(header + 1, ws.max_row + 1):
            comment = normalize(ws.cell(r, columns["注释"]).value)
            if comment:
                lookup[comment].add((clean_text(ws.cell(r, columns["地址"]).value), clean_text(ws.cell(r, columns["数据类型"]).value)))
    finally:
        wb.close()
    counts = Counter(make_initials(name) for name, _ in parameters)
    rows, warnings, used = [], [], set()
    matched = 0
    for name, unit in parameters:
        identifier = (expand_number_initials(name) if counts[make_initials(name)] > 1 else make_initials(name)).lower()
        if identifier[0].isdigit(): identifier = "p_" + identifier
        if identifier in used:
            raise ValueError(f"点位标识重复：{name} → {identifier}")
        used.add(identifier)
        address, data_type = "", ""
        for candidate in match_names(name, dual):
            records = lookup.get(normalize(candidate))
            if not records:
                continue
            if len(records) != 1:
                warnings.append(f"{name}：注释 {candidate} 对应多个不同点位，地址和类型留空")
                break
            raw_address, raw_type = next(iter(records))
            if modbus:
                address, correction = convert_modbus_address(raw_address, candidate, device_file)
                if correction:
                    warnings.append(correction)
            else:
                address = extract_address(raw_address)
            data_type = convert_type(raw_type)
            if not address: warnings.append(f"{name}：无法解析地址 {raw_address}")
            if not data_type: warnings.append(f"{name}：不支持的数据类型 {raw_type}")
            if address and data_type: matched += 1
            break
        else:
            warnings.append(f"{name}：未找到对应注释，地址和类型留空")
        rows.append((name, identifier, address, data_type, unit, "", 1000, "", "", ""))
    output_file = output_dir / f"{parameter_file.stem}_{device_file.stem}_点位导入.xlsx"
    if output_file in (parameter_file, device_file, template_file):
        raise ValueError("输出文件不能覆盖输入文件。")
    wb = load_workbook(template_file)
    temporary = None
    try:
        ws = wb.worksheets[0]
        header = find_header_row(ws, POINT_COLUMNS)
        columns = get_header_map(ws, header)
        for row in ws.iter_rows(min_row=header + 1):
            for cell in row: cell.value = None
        for r, values in enumerate(rows, header + 1):
            if r > header + 1: copy_row_style(ws, header + 1, r)
            for field, value in zip(POINT_COLUMNS, values):
                cell = ws.cell(r, columns[field], value)
                if isinstance(value, str): cell.data_type = "s"
        last = header + len(rows)
        if ws.max_row > last: ws.delete_rows(last + 1, ws.max_row - last)
        output_dir.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=output_dir, suffix=".xlsx", delete=False) as temp:
            temporary = Path(temp.name)
        wb.save(temporary)
        os.replace(temporary, output_file)
    finally:
        wb.close()
        if temporary and temporary.exists(): temporary.unlink()
    if progress: progress(1, 1, parameter_file.stem, output_file)
    return {"generated_files": [output_file], "output_dir": output_dir, "warnings": warnings,
            "point_count": len(rows), "matched_count": matched, "rule": ("双色" if dual else "单色") + (" / Modbus" if modbus else " / OPC")}

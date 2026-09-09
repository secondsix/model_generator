import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from point_generator import POINT_COLUMNS, POINT_TYPES, match_names, convert_type, extract_address, generate_point_file
from model_generator_gui import ModelGeneratorApp


class PointTests(unittest.TestCase):
    def test_single_fallback_priority(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, device, template = root / "单色工艺参数.xlsx", root / "device.xlsx", root / "template.xlsx"
            fixtures = [(source, [["工艺参数名称", "工艺参数单位"], ["注射压力一段", "bar"], ["注射压力二段", "bar"], ["切换时间", "s"], ["射嘴温度", "℃"]]),
                        (device, [["注释", "地址", "数据类型"], ["射出1段压力", "ns=2;i=1", "INT32"], ["注射压力1", "ns=2;i=9", "FLOAT32"], ["注射压力2", "ns=2;i=2", "FLOAT32"], ["切保压时间（射胶时间）", "ns=2;i=3", "FLOAT32"]]),
                        (template, [list(POINT_COLUMNS)])]
            for path, rows in fixtures:
                wb = Workbook()
                for row in rows: wb.active.append(row)
                wb.save(path)
                wb.close()
            result = generate_point_file(source, device, template, root / "output")
            wb = load_workbook(result["generated_files"][0])
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual([row[2] for row in rows], ["ns=2;i=1", "ns=2;i=2", "ns=2;i=3", None])
            self.assertEqual(result["matched_count"], 3)
        self.assertNotIn("注射压力1", match_names("A组注射压力一段", True))
        for stem, count in (("注射压力", 4), ("保压压力", 4), ("注射速度", 6), ("保压时间", 4), ("储料位置", 3)):
            for index, number in enumerate("一二三四五六"[:count], 1):
                self.assertIn(f"{stem}{index}", match_names(f"{stem}{number}段"))
        for unit, index in (("mm/s", 1), ("r/min", 2), ("%", 3)):
            self.assertIn(f"储料速度{index}", match_names(f"储料速度（{unit}）"))

    def test_third_rule_priority_and_export(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, device, template = root / "单色工艺参数.xlsx", root / "device.xlsx", root / "template.xlsx"
            names = ["注射速度一段（mm/s）", "注射速度二段（%）", "注射速度三段（mm/s）", "切换时间", "余料位置", "注射速度四段（%）"]
            fixtures = [(source, [["工艺参数名称", "工艺参数单位"]] + [[name, ""] for name in names]),
                        (device, [["注释", "地址", "数据类型"], ["射出1段速度", "ns=2;i=1", "INT32"], ["注射速度1", "ns=2;i=11", "INT32"], ["注射速度1段", "ns=2;i=12", "INT32"], ["注射速度2", "ns=2;i=2", "INT32"], ["注射速度2段", "ns=2;i=22", "INT32"], ["注射速度3段", "ns=2;i=3", "FLOAT32"], ["注射切换时间", "ns=2;i=4", "FLOAT32"], ["余料位置", "ns=2;i=5", "FLOAT32"], ["注射速度4段", "ns=2;i=6", "INT32"]]),
                        (template, [list(POINT_COLUMNS)])]
            for path, rows in fixtures:
                wb = Workbook()
                for row in rows: wb.active.append(row)
                wb.save(path)
                wb.close()
            result = generate_point_file(source, device, template, root / "output")
            wb = load_workbook(result["generated_files"][0])
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual([row[2] for row in rows], ["ns=2;i=1", "ns=2;i=2", "ns=2;i=3", "ns=2;i=4", "ns=2;i=5", None])
        for index, number in enumerate("一二三", 1):
            for unit in ("mm/s", "%"):
                candidates = match_names(f"注射速度{number}段（{unit}）")
                self.assertLess(candidates.index(f"注射速度{index}"), candidates.index(f"注射速度{index}段"))
        self.assertNotIn("注射速度1段", match_names("A组注射速度一段", True))

    def test_fourth_rule_priority_and_export(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, device, template = root / "单色工艺参数.xlsx", root / "device.xlsx", root / "template.xlsx"
            names = ["注射速度一段（%）", "注射速度二段（%）", "注射速度二段（mm/s）", "保压压力四段", "切换保压位置", "注射终点位置"]
            fixtures = [(source, [["工艺参数名称", "工艺参数单位"]] + [[name, ""] for name in names]),
                        (device, [["注释", "地址", "数据类型"], ["射出1段速度", "ns=2;i=1", "INT32"], ["射出1段速度B", "ns=2;i=11", "FLOAT32"], ["射出2段速度B", "ns=2;i=2", "FLOAT32"], ["保压4段压力", "ns=2;i=4", "FLOAT32"], ["转保压位置", "ns=2;i=5", "FLOAT32"], ["射出终点", "ns=2;i=6", "FLOAT32"]]),
                        (template, [list(POINT_COLUMNS)])]
            for path, rows in fixtures:
                wb = Workbook()
                for row in rows: wb.active.append(row)
                wb.save(path)
                wb.close()
            result = generate_point_file(source, device, template, root / "output")
            wb = load_workbook(result["generated_files"][0])
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual([row[2] for row in rows], ["ns=2;i=1", "ns=2;i=2", None, "ns=2;i=4", "ns=2;i=5", "ns=2;i=6"])
        for index, number in enumerate("一二三四五六", 1):
            candidates = match_names(f"注射速度{number}段（%）")
            self.assertLess(candidates.index(f"注射速度{index}"), candidates.index(f"射出{index}段速度B"))
            self.assertNotIn(f"射出{index}段速度B", match_names(f"注射速度{number}段（mm/s）"))
            self.assertNotIn(f"射出{index}段速度B", match_names(f"A组注射速度{number}段", True))

    def test_dual_fallback_priority_and_special_names(self):
        for group in "AB":
            for stem, count in (("注射压力", 4), ("保压压力", 4), ("注射速度", 6), ("保压时间", 4), ("储料位置", 3)):
                for index, number in enumerate("一二三四五六"[:count], 1):
                    candidates = match_names(f"{group}组{stem}{number}段", True)
                    self.assertEqual(candidates[-1], f"{group}组{stem}{index}")
                    other = "B" if group == "A" else "A"
                    self.assertNotIn(f"{other}组{stem}{index}", candidates)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, device, template = root / "双色工艺参数.xlsx", root / "device.xlsx", root / "template.xlsx"
            names = ["A组注射压力一段", "B组注射压力一段", "A组加料时间", "B组加料时间", "A组射出终点位置", "B组射出终点位置", "B组切换时间", "A组机筒温度一段"]
            fixtures = [(source, [["工艺参数名称", "工艺参数单位"]] + [[name, ""] for name in names]),
                        (device, [["注释", "地址", "数据类型"], ["射出一段压力A", "ns=2;i=1", "INT32"], ["A组注射压力1", "ns=2;i=11", "INT32"], ["B组注射压力1", "ns=2;i=2", "FLOAT32"], ["储料时间", "ns=2;i=3", "FLOAT32"], ["B储料时间", "ns=2;i=4", "FLOAT32"], ["射出终点位置", "ns=2;i=5", "FLOAT32"], ["B组切保压时间（射胶时间）", "ns=2;i=6", "FLOAT32"]]),
                        (template, [list(POINT_COLUMNS)])]
            for path, rows in fixtures:
                wb = Workbook()
                for row in rows: wb.active.append(row)
                wb.save(path)
                wb.close()
            result = generate_point_file(source, device, template, root / "output")
            wb = load_workbook(result["generated_files"][0])
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual([row[2] for row in rows], ["ns=2;i=1", "ns=2;i=2", "ns=2;i=3", "ns=2;i=4", "ns=2;i=5", "ns=2;i=5", "ns=2;i=6", None])
        self.assertNotIn("A组注射压力1", match_names("A组注射压力一段", False))

    def test_fifth_rule_export_and_priority(self):
        for index, number in enumerate("一二三四五六"):
            target = "设定值注入速度" + (str(index) if index else "")
            for unit in ("mm/s", "%"):
                candidates = match_names(f"注射速度{number}段（{unit}）")
                self.assertIn(target, candidates)
                self.assertNotIn(target, match_names(f"A组注射速度{number}段", True))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, device, template = root / "单色工艺参数.xlsx", root / "device.xlsx", root / "template.xlsx"
            names = ["注射速度一段（mm/s）", "注射速度二段（%）", "背压", "射嘴温度", "喷嘴温度", "切换时间", "保压起点"]
            fixtures = [(source, [["工艺参数名称", "工艺参数单位"]] + [[name, ""] for name in names]),
                        (device, [["注释", "地址", "数据类型"], ["射出1段速度", "ns=2;i=1", "INT32"], ["设定值注入速度", "ns=2;i=11", "FLOAT32"], ["设定值注入速度1", "ns=2;i=2", "FLOAT32"], ["ZS450B_背压设定值", "ns=2;i=3", "FLOAT32"], ["ZS450B_射嘴实际温度", "ns=2;i=4", "FLOAT32"], ["实际值温度喷嘴1", "ns=2;i=5", "FLOAT32"], ["实际值PN切换时间", "ns=2;i=6", "FLOAT32"]]),
                        (template, [list(POINT_COLUMNS)])]
            for path, rows in fixtures:
                wb = Workbook()
                for row in rows: wb.active.append(row)
                wb.save(path)
                wb.close()
            result = generate_point_file(source, device, template, root / "output")
            wb = load_workbook(result["generated_files"][0])
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual([row[2] for row in rows], ["ns=2;i=1", "ns=2;i=2", "ns=2;i=3", "ns=2;i=4", "ns=2;i=5", "ns=2;i=6", None])

    def test_sixth_rule_export_and_scope(self):
        expected = {"注射压力一段": "ZS474双色_注射压力", "背压": "ZS474双色_背压设定值", "注射速度一段（mm/s）": "ZS474双色_注射速度", "注射时间": "ZS474双色_注射时间", "螺杆转速(r/min)": "ZS474双色_螺杆转速设定", "螺杆转速(%)": "ZS474双色_螺杆转速设定", "冷却时间": "ZS474双色_冷却时间设定值", "加料时间": "ZS474双色_加料时间", "切换保压位置": "ZS474双色_切换位置实际值", "余料位置": "ZS474双色_余料位置", "射嘴温度": "ZS474SS_射嘴实际温度"}
        for name, target in expected.items():
            self.assertIn(target, match_names(name))
            self.assertNotIn(target, match_names("A组" + name, True))
        for name in ("注射速度一段（%）", "注射速度二段（mm/s）", "喷嘴温度", "注射压力二段"):
            self.assertFalse(any(value.startswith("ZS474") for value in match_names(name)))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, device, template = root / "单色工艺参数.xlsx", root / "device.xlsx", root / "template.xlsx"
            fixtures = [(source, [["工艺参数名称", "工艺参数单位"], ["背压", "bar"], ["注射压力一段", "bar"], ["注射速度一段（%）", "%"], ["射嘴温度", "℃"]]),
                        (device, [["注释", "地址", "数据类型"], ["ZS450B_背压设定值", "ns=2;i=1", "INT32"], ["ZS474双色_背压设定值", "ns=2;i=11", "FLOAT32"], ["ZS474双色_注射压力", "ns=2;i=2", "FLOAT32"], ["ZS474双色_注射速度", "ns=2;i=3", "FLOAT32"], ["ZS474SS_射嘴实际温度", "ns=2;i=4", "FLOAT32"]]),
                        (template, [list(POINT_COLUMNS)])]
            for path, rows in fixtures:
                wb = Workbook()
                for row in rows: wb.active.append(row)
                wb.save(path)
                wb.close()
            result = generate_point_file(source, device, template, root / "output")
            wb = load_workbook(result["generated_files"][0])
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual([row[2] for row in rows], ["ns=2;i=1", "ns=2;i=2", None, "ns=2;i=4"])

    def test_seventh_rule_export_and_priority(self):
        expected = {"注射压力一段": "注射压力", "保压压力一段": "保压压力", "背压": "背压设定值", "注射速度一段（mm/s）": "注射速度", "注射速度一段（%）": "注射速度", "注射时间": "注射时间", "螺杆转速(r/min)": "螺杆转速设定", "螺杆转速(%)": "螺杆转速设定", "保压时间一段": "保压时间", "冷却时间": "冷却时间设定值", "加料时间": "加料时间", "切换保压位置": "切换位置", "余料位置": "余料位置", "射嘴温度": "射嘴实际温度"}
        for name, suffix in expected.items():
            target = "ZS451双色_" + suffix
            self.assertIn(target, match_names(name))
            self.assertNotIn(target, match_names("A组" + name, True))
        for name in ("注射压力二段", "保压压力二段", "保压时间二段", "注射速度二段（%）", "喷嘴温度"):
            self.assertFalse(any(v.startswith("ZS451") for v in match_names(name)))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, device, template = root / "单色工艺参数.xlsx", root / "device.xlsx", root / "template.xlsx"
            names = ["注射压力一段", "保压压力一段", "注射速度一段（mm/s）", "注射速度一段（%）", "保压时间一段", "切换保压位置", "喷嘴温度"]
            fixtures = [(source, [["工艺参数名称", "工艺参数单位"]] + [[name, ""] for name in names]),
                        (device, [["注释", "地址", "数据类型"], ["ZS474双色_注射压力", "ns=2;i=1", "INT32"], ["ZS451双色_注射压力", "ns=2;i=11", "FLOAT32"], ["ZS451双色_保压压力", "ns=2;i=2", "FLOAT32"], ["ZS451双色_注射速度", "ns=2;i=3", "FLOAT32"], ["ZS451双色_保压时间", "ns=2;i=4", "FLOAT32"], ["ZS451双色_切换位置", "ns=2;i=5", "FLOAT32"]]),
                        (template, [list(POINT_COLUMNS)])]
            for path, rows in fixtures:
                wb = Workbook()
                for row in rows: wb.active.append(row)
                wb.save(path)
                wb.close()
            result = generate_point_file(source, device, template, root / "output")
            wb = load_workbook(result["generated_files"][0])
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual([row[2] for row in rows], ["ns=2;i=1", "ns=2;i=2", "ns=2;i=3", "ns=2;i=3", "ns=2;i=4", "ns=2;i=5", None])

    def test_eighth_rule_export_and_scope(self):
        expected = {"冷却时间": "实际冷却时间", "射嘴温度": "实际射嘴温度"}
        for index, number in enumerate("一二三四", 1):
            expected[f"注射压力{number}段"] = f"设定射胶{index}段压力"
            expected[f"保压压力{number}段"] = f"设定保压{index}段压力"
            expected[f"保压时间{number}段"] = f"设定保压{index}段时间"
        for index, number in enumerate("一二三四五", 1):
            for unit in ("mm/s", "%"):
                expected[f"注射速度{number}段（{unit}）"] = f"设定射胶{index}段速度"
        for name, target in expected.items():
            self.assertEqual(match_names(name)[-1], target)
            self.assertNotIn(target, match_names("A组" + name, True))
        for unit in ("mm/s", "%"):
            self.assertNotIn("设定射胶6段速度", match_names(f"注射速度六段（{unit}）"))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, device, template = root / "单色工艺参数.xlsx", root / "device.xlsx", root / "template.xlsx"
            names = ["注射压力一段", "保压压力四段", "保压时间四段", "注射速度五段（%）", "冷却时间", "射嘴温度", "注射速度六段（mm/s）"]
            fixtures = [(source, [["工艺参数名称", "工艺参数单位"]] + [[name, ""] for name in names]),
                        (device, [["注释", "地址", "数据类型"], ["ZS451双色_注射压力", "ns=2;i=1", "INT32"], ["设定射胶1段压力", "ns=2;i=11", "FLOAT32"], ["设定保压4段压力", "ns=2;i=2", "FLOAT32"], ["设定保压4段时间", "ns=2;i=3", "FLOAT32"], ["设定射胶5段速度", "ns=2;i=4", "FLOAT32"], ["实际冷却时间", "ns=2;i=5", "FLOAT32"], ["实际射嘴温度", "ns=2;i=6", "FLOAT32"], ["设定射胶6段速度", "ns=2;i=7", "FLOAT32"]]),
                        (template, [list(POINT_COLUMNS)])]
            for path, rows in fixtures:
                wb = Workbook()
                for row in rows: wb.active.append(row)
                wb.save(path)
                wb.close()
            result = generate_point_file(source, device, template, root / "output")
            wb = load_workbook(result["generated_files"][0])
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual([row[2] for row in rows], ["ns=2;i=1", "ns=2;i=2", "ns=2;i=3", "ns=2;i=4", "ns=2;i=5", "ns=2;i=6", None])

    def test_mapping_and_types(self):
        self.assertEqual(match_names("注射压力三段")[0], "射出3段压力")
        self.assertEqual(match_names("保压压力四段")[0], "保压第段压力4")
        self.assertEqual(match_names("储料速度（r/min）")[0], "储料2段速度")
        for group in "AB":
            self.assertEqual(match_names(f"{group}组注射压力三段", True)[0], f"射出三段压力{group}")
            self.assertEqual(match_names(f"{group}组保压压力四段", True)[0], f"保压第四段压力{group}")
            self.assertEqual(match_names(f"{group}组冷却时间", True)[0], f"冷却实际时间{group}")
        self.assertEqual(convert_type("有符号的 32 位值"), "INT32")
        self.assertEqual(convert_type("无符号的 16 位值"), "UINT16")
        self.assertEqual(convert_type("32-位浮点数 IEEE 754"), "FLOAT32")
        self.assertEqual(convert_type("64-位浮点数 IEEE 754"), "FLOAT64")
        for kind in POINT_TYPES:
            self.assertEqual(convert_type(kind), kind)
        self.assertEqual(convert_type("未知类型"), "")
        self.assertEqual(extract_address('"ns=2;i=20250","6;-1",0'), "ns=2;i=20250")

    def test_dual_export_and_missing(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, device, template = root / "双色注塑机工艺参数.xlsx", root / "device.xlsx", root / "template.xlsx"
            fixtures = [(source, [["工艺参数名称", "工艺参数单位"], ["A组注射压力一段", "bar"], ["B组注射压力一段", "bar"], ["A组机筒温度七段", "℃"], ["A组冷却时间", "s"]]),
                        (device, [["说明行"], ["注释", "地址", "数据类型"], ["射出一段压力A", '"ns=2;i=1","6;-1",0', "有符号的 32 位值"], ["射出一段压力B", '"ns=2;i=2","6;-1",0', "32-位浮点数 IEEE 754"], ["七段温度异常报警A", "ns=2;i=3", "BOOL"], ["冷却实际时间A", "ns=2;i=4", "INT32"], ["冷却实际时间A", "ns=2;i=5", "INT32"]]),
                        (template, [list(POINT_COLUMNS), ["示例", "old", "HR0", "UINT16", "℃", "旧标签", 500, "旧配置", "旧类型", "旧转换"]])]
            for path, rows in fixtures:
                wb = Workbook()
                for row in rows: wb.active.append(row)
                wb.save(path)
                wb.close()
            result = generate_point_file(source, device, template, root / "output")
            self.assertEqual(result["matched_count"], 2)
            self.assertEqual(len(result["warnings"]), 2)
            wb = load_workbook(result["generated_files"][0])
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual(rows[0][:5], ("A组注射压力一段", "azzsylyd", "ns=2;i=1", "INT32", "bar"))
            self.assertEqual(rows[1][2:4], ("ns=2;i=2", "FLOAT32"))
            self.assertEqual(rows[2][2:4], (None, None))
            self.assertEqual(rows[3][2:4], (None, None))
            for row in rows:
                self.assertEqual(row[5:], (None, 1000, None, None, None))

    def test_ui_point_mode(self):
        app = ModelGeneratorApp()
        app.withdraw()
        try:
            app.mode_var.set("单色注塑机点位导入生成器")
            app._switch_mode()
            self.assertTrue(app._single_source())
            self.assertEqual(app.source_label.cget("text"), "工艺参数")
            with patch("model_generator_gui.filedialog.askopenfilename", return_value="C:/单色工艺参数.xlsx"):
                app._choose_source()
            self.assertEqual(app.source_var.get(), "C:/单色工艺参数.xlsx")
            app.device_var.set("C:/device.xlsx")
            app.mode_var.set("物模型标签库导入生成器")
            app._switch_mode()
            app.mode_var.set("单色注塑机点位导入生成器")
            app._switch_mode()
            self.assertEqual(app.device_var.get(), "C:/device.xlsx")
        finally:
            app.destroy()


if __name__ == "__main__":
    unittest.main()

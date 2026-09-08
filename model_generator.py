"""根据物模型设计表批量生成平台导入文件。"""

import argparse
import json
import re
import os
import tempfile
from collections import Counter
from copy import copy
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from pypinyin import lazy_pinyin
except ModuleNotFoundError as exc:
    raise SystemExit(
        "缺少运行依赖，请先执行：pip install openpyxl pypinyin"
    ) from exc


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_OUTPUT_DIR = SCRIPT_DIR / "output"

# 设计表第一个 sheet 的实际字段。
SOURCE_COLUMNS = {
    "product_name": "产品名称",
    "model_type": "物模型类型",
    "property_name": "物模型名称",
    "unit": "单位",
    "description": "备注说明",
}

# 模板第一个 sheet 的实际字段。
TEMPLATE_COLUMNS = (
    "属性标识",
    "属性名称",
    "数据类型",
    "单位",
    "精度",
    "数据类型配置",
    "来源",
    "属性说明",
    "读写类型",
    "存储方式",
)


def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def make_identifier(name: str) -> str:
    """把中文属性名称转换为适合导入的拼音属性标识。"""
    parts = lazy_pinyin(clean_text(name))
    identifier = "_".join(parts)
    identifier = re.sub(r"[^a-zA-Z0-9_]", "_", identifier)
    identifier = re.sub(r"_+", "_", identifier).strip("_").lower()

    if not identifier:
        identifier = "property"
    if identifier[0].isdigit():
        identifier = f"p_{identifier}"
    return identifier


def make_unique_identifier(name: str, used_identifiers: Counter) -> str:
    """保证同一个产品内的属性标识不重复。"""
    base = make_identifier(name)
    used_identifiers[base] += 1
    sequence = used_identifiers[base]
    return base if sequence == 1 else f"{base}_{sequence}"


def safe_filename(name: str) -> str:
    """替换 Windows 文件名中的非法字符。"""
    result = re.sub(r'[<>:"/\\|?*]', "_", clean_text(name)).rstrip(". ")
    return result or "未命名物模型"


def expand_property_name(name: str):
    """将“A/B组……”或“A／B组……”拆成 A、B 两个属性。"""
    name = clean_text(name)
    for marker in ("A/B组", "A／B组"):
        if marker in name:
            return [name.replace(marker, "A组"), name.replace(marker, "B组")]
    return [name]


def is_parameter_type(model_type: str) -> bool:
    """当前仅处理名称以“参数”结尾的物模型类型。"""
    return clean_text(model_type).endswith("参数")


ALLOWED_DATA_TYPES = {"boolean", "date", "double", "enum", "float", "int", "long", "string", "geoPoint", "file", "password"}


def get_data_type(name: str, unit: str = ""):
    """单位优先；无单位时按名称辅助判断，未知内容作为字符串。"""
    name, unit = clean_text(name), clean_text(unit)
    if unit in ALLOWED_DATA_TYPES:
        kind, output_unit = unit, ""
    elif unit in {"次", "个", "件", "台", "只", "count", "pcs"}:
        kind = "long" if re.search(r"累计|总", name) else "int"
        output_unit = unit
    elif unit in {"是/否", "开/关", "true/false"}:
        kind, output_unit = "boolean", ""
    elif unit in {"日期", "时间戳", "timestamp"}:
        kind, output_unit = "date", ""
    elif unit and unit not in {"-", "—", "/", "无", "无单位"}:
        kind, output_unit = "double", unit
    else:
        output_unit = ""
        if re.search(r"密码|口令", name):
            kind = "password"
        elif re.search(r"文件|图片|附件", name):
            kind = "file"
        elif re.search(r"经纬度|地理坐标", name):
            kind = "geoPoint"
        elif re.search(r"日期|时刻|时间戳", name):
            kind = "date"
        elif re.search(r"是否|故障|报警|开关|启停", name):
            kind = "boolean"
        elif re.search(r"状态|模式|枚举", name):
            kind = "enum"
        elif re.search(r"总产量|累计产量|累计次数|总次数", name):
            kind = "long"
        elif re.search(r"次数|数量|个数|计数", name):
            kind = "int"
        else:
            kind = "string"
    config = {"type": kind}
    precision = ""
    if kind in {"double", "float", "int", "long"}:
        config["round"] = "HALF_UP"
        if output_unit:
            config["unit"] = output_unit
        if kind in {"double", "float"}:
            precision = 2
            config["scale"] = precision
    elif kind == "boolean":
        config.update(falseValue="false", trueValue="true", trueText="是", falseText="否")
    elif kind == "date":
        config.update(tz="Asia/Shanghai", format="timestamp")
    elif kind == "geoPoint":
        config.update(lonProperty="lon", latProperty="lat")
    elif kind == "file":
        config.update(bodyType="url", mediaType="*/*")
    return {"data_type": kind, "unit": output_unit, "precision": precision,
            "config": json.dumps(config, ensure_ascii=False, separators=(",", ":"))}


def read_tag_library(source_file):
    """读取标签库首个 sheet，原样保留标签标识和名称。"""
    wb = load_workbook(source_file, data_only=True)
    try:
        ws = wb.worksheets[0]
        header = find_header_row(ws, ("标识", "名称", "单位"))
        headers = get_header_map(ws, header)
        parameters, identifiers = [], set()
        skipped = Counter()
        for row in range(header + 1, ws.max_row + 1):
            def value(field):
                return clean_text(ws.cell(row, headers[field]).value) if field in headers else ""
            identifier, name = value("标识"), value("名称")
            if not identifier and not name:
                continue
            if value("标签类型") not in ("", "property"):
                skipped[value("标签类型")] += 1
                continue
            if not identifier or not name:
                raise ValueError(f"第 {row} 行缺少标识或名称")
            if identifier in identifiers:
                raise ValueError(f"第 {row} 行属性标识重复：{identifier}")
            identifiers.add(identifier)
            parameters.append({"identifier": identifier, "name": name, "unit": value("单位"), "description": value("说明")})
        if not parameters:
            raise ValueError("标签库中没有可导出的 property 标签。")
        product = source_file.stem.removesuffix("物模型标签库导入") or "物模型"
        return ws.title, {product: {"parameters": parameters}}, Counter(property=len(parameters)), skipped
    finally:
        wb.close()


def find_header_row(ws, required_columns, max_rows=30):
    """在工作表前若干行中查找同时包含指定字段的表头。"""
    required = set(required_columns)
    for row in range(1, min(ws.max_row, max_rows) + 1):
        values = {
            clean_text(ws.cell(row=row, column=col).value)
            for col in range(1, ws.max_column + 1)
        }
        if required.issubset(values):
            return row
    raise RuntimeError(f"找不到表头，需要包含：{', '.join(required_columns)}")


def get_header_map(ws, header_row):
    """返回“表头名称 -> Excel 列号”的映射。"""
    result = {}
    for col in range(1, ws.max_column + 1):
        value = clean_text(ws.cell(row=header_row, column=col).value)
        if value:
            result[value] = col
    return result


def get_optional_cell(ws, row, headers, field_name):
    column_name = SOURCE_COLUMNS[field_name]
    column = headers.get(column_name)
    return "" if column is None else clean_text(ws.cell(row, column).value)


def copy_row_style(ws, source_row, target_row):
    """复制模板数据行的格式和行高。"""
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.protection = copy(source.protection)


def display_width(value) -> int:
    """粗略计算 Excel 中英文混排文本所需宽度。"""
    text = clean_text(value)
    return sum(2 if ord(char) > 127 else 1 for char in text)


def adjust_output_layout(ws, header_row, last_row):
    """按生成内容调整列宽，并让长标识、配置和说明自动换行。"""
    min_widths = {1: 20, 2: 16, 3: 10, 4: 10, 5: 8, 6: 18, 7: 10, 8: 18, 9: 10, 10: 10}
    max_widths = {1: 45, 2: 36, 3: 14, 4: 16, 5: 10, 6: 60, 7: 12, 8: 60, 9: 14, 10: 14}

    actual_widths = {}
    for col in range(1, ws.max_column + 1):
        content_width = max(
            display_width(ws.cell(row, col).value)
            for row in range(header_row, last_row + 1)
        ) + 2
        width = max(min_widths.get(col, 10), content_width)
        width = min(max_widths.get(col, 40), width)
        actual_widths[col] = width
        ws.column_dimensions[get_column_letter(col)].width = width

    wrap_columns = (1, 2, 6, 8)
    for row in range(header_row + 1, last_row + 1):
        line_count = 1
        for col in wrap_columns:
            cell = ws.cell(row, col)
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = "top"
            cell.alignment = alignment
            width = max(actual_widths[col] - 2, 1)
            line_count = max(line_count, (display_width(cell.value) + int(width) - 1) // int(width))
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, line_count * 15)


def read_models(source_file: Path, expand_names=True):
    """读取设计表第一个 sheet，并按产品名称汇总“xx参数”属性。"""
    wb = load_workbook(source_file, data_only=True)
    try:
        if not wb.worksheets:
            raise RuntimeError(f"{source_file.name} 不包含工作表")

        # 用户要求：按索引读取第一个 sheet，不依赖 sheet 名称。
        ws = wb.worksheets[0]
        required_headers = (
            SOURCE_COLUMNS["product_name"],
            SOURCE_COLUMNS["model_type"],
            SOURCE_COLUMNS["property_name"],
        )
        header_row = find_header_row(ws, required_headers)
        headers = get_header_map(ws, header_row)

        models = {}
        skipped_types = Counter()
        included_types = Counter()
        current_product_name = ""

        for row in range(header_row + 1, ws.max_row + 1):
            # “产品名称”使用了纵向合并单元格；空白格继承上一非空产品名称。
            product_value = clean_text(
                ws.cell(row, headers[SOURCE_COLUMNS["product_name"]]).value
            )
            if product_value:
                current_product_name = product_value

            model_type = clean_text(
                ws.cell(row, headers[SOURCE_COLUMNS["model_type"]]).value
            )
            property_name = clean_text(
                ws.cell(row, headers[SOURCE_COLUMNS["property_name"]]).value
            )

            if not model_type or not property_name:
                continue
            if not is_parameter_type(model_type):
                skipped_types[model_type] += 1
                continue
            if not current_product_name:
                raise RuntimeError(f"第 {row} 行缺少产品名称，且无法从上一行继承")

            model = models.setdefault(current_product_name, {"parameters": []})
            unit = get_optional_cell(ws, row, headers, "unit")
            description = get_optional_cell(ws, row, headers, "description")

            for expanded_name in (expand_property_name(property_name) if expand_names else [property_name]):
                model["parameters"].append(
                    {
                        "name": expanded_name,
                        "unit": unit,
                        "description": (description or expanded_name) if expand_names else description,
                        "model_type": model_type,
                    }
                )
                included_types[model_type] += 1

        return ws.title, models, included_types, skipped_types
    finally:
        wb.close()


def generate_model_file(model_name, model_info, template_file: Path, output_dir: Path):
    """复制实际导入模板，并写入一个产品的全部参数属性。"""
    output_dir.mkdir(parents=True, exist_ok=True)
    output_file = output_dir / f"{safe_filename(model_name)}物模型导入.xlsx"
    if output_file.resolve() == template_file.resolve():
        raise ValueError("保存位置不能与模板相同。")

    wb = load_workbook(template_file)
    temporary = None
    try:
        if not wb.worksheets:
            raise RuntimeError(f"{template_file.name} 不包含工作表")
        ws = wb.worksheets[0]
        header_row = find_header_row(ws, TEMPLATE_COLUMNS)
        headers = get_header_map(ws, header_row)
        data_start_row = header_row + 1
        style_row = data_start_row

        # 清除模板示例值，但保留第一条示例行的格式作为生成行样式。
        original_max_row = ws.max_row
        for row in range(data_start_row, original_max_row + 1):
            for col in range(1, ws.max_column + 1):
                ws.cell(row, col).value = None

        current_row = data_start_row
        for parameter in model_info["parameters"]:
            if current_row != style_row:
                copy_row_style(ws, style_row, current_row)

            name = parameter["name"]
            type_info = get_data_type(name, parameter["unit"])
            values = {
                "属性标识": parameter["identifier"],
                "属性名称": name,
                "数据类型": type_info["data_type"],
                "单位": type_info["unit"],
                "精度": type_info["precision"],
                "数据类型配置": type_info["config"],
                "来源": "设备",
                "属性说明": parameter["description"],
                "读写类型": "上报",
                "存储方式": "存储",
            }
            for field, value in values.items():
                cell = ws.cell(current_row, headers[field], value)
                if isinstance(value, str):
                    cell.data_type = "s"
            current_row += 1

        # 删除未被覆盖的模板示例行，避免示例数据混入导入结果。
        if current_row <= ws.max_row:
            ws.delete_rows(current_row, ws.max_row - current_row + 1)

        adjust_output_layout(ws, header_row, current_row - 1)

        with tempfile.NamedTemporaryFile(dir=output_dir, suffix=".xlsx", delete=False) as temp:
            temporary = Path(temp.name)
        wb.save(temporary)
        os.replace(temporary, output_file)
    finally:
        wb.close()
        if temporary and temporary.exists():
            temporary.unlink()

    print(f"[生成成功] {output_file.name}  属性数：{len(model_info['parameters'])}")
    return output_file


def parse_args():
    parser = argparse.ArgumentParser(description="Excel 物模型批量生成工具")
    parser.add_argument("--source", type=Path, required=True, help="物模型标签库导入文件路径")
    parser.add_argument("--template", type=Path, required=True, help="物模型导入模板路径")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR, help="输出目录")
    parser.add_argument("--only", help="只生成指定产品名称，便于测试")
    return parser.parse_args()


def generate_files(source_file: Path, template_file: Path, output_dir: Path, only=None, progress=None):
    """执行批量生成，并返回可供命令行或图形界面展示的结果。"""
    source_file = Path(source_file).resolve()
    template_file = Path(template_file).resolve()
    output_dir = Path(output_dir).resolve()
    if not source_file.exists():
        raise FileNotFoundError(f"找不到文件：{source_file}")
    if not template_file.exists():
        raise FileNotFoundError(f"找不到文件：{template_file}")

    sheet_name, models, included_types, skipped_types = read_tag_library(source_file)
    if only:
        if only not in models:
            available = "、".join(models.keys())
            raise RuntimeError(f"找不到产品：{only}。可选产品：{available}")
        models = {only: models[only]}

    generated_files = []
    total = len(models)
    for index, (model_name, model_info) in enumerate(models.items(), start=1):
        if (output_dir / f"{safe_filename(model_name)}物模型导入.xlsx").resolve() == source_file:
            raise ValueError("保存位置不能与标签库源文件相同。")
        generated_files.append(
            generate_model_file(model_name, model_info, template_file, output_dir)
        )
        if progress:
            progress(index, total, model_name, generated_files[-1])

    return {
        "sheet_name": sheet_name,
        "included_types": included_types,
        "skipped_types": skipped_types,
        "generated_files": generated_files,
        "output_dir": output_dir,
    }


def generate_many_files(source_files, template_file, output_dir, progress=None):
    """预检所有标签库，再逐个生成，防止同名文件相互覆盖。"""
    sources = list(dict.fromkeys(Path(path).resolve() for path in source_files))
    if not sources:
        raise ValueError("请至少选择一个标签库文件。")
    template_file, output_dir = Path(template_file).resolve(), Path(output_dir).resolve()
    planned, destinations = [], set()
    for source in sources:
        _, models, _, _ = read_tag_library(source)
        for name, model in models.items():
            destination = (output_dir / f"{safe_filename(name)}物模型导入.xlsx").resolve()
            if destination in destinations:
                raise ValueError(f"所选标签库的输出文件名重复：{destination.name}")
            if destination in sources or destination == template_file:
                raise ValueError(f"输出文件不能覆盖源文件或模板：{destination}")
            destinations.add(destination)
            planned.append((name, model))
    generated = []
    for index, (name, model) in enumerate(planned, 1):
        generated.append(generate_model_file(name, model, template_file, output_dir))
        if progress:
            progress(index, len(planned), name, generated[-1])
    return {"generated_files": generated, "output_dir": output_dir}


def main():
    args = parse_args()
    result = generate_files(
        args.source,
        args.template,
        args.output_dir,
        only=args.only,
    )

    print("=" * 60)
    print("Excel 物模型批量生成工具")
    print(f"标签库第一个 sheet：{result['sheet_name']}")
    print("处理规则：从标签库读取 property 标签，按单位判断数据类型")
    print(f"纳入类型：{dict(result['included_types'])}")
    print(f"跳过类型：{dict(result['skipped_types'])}")
    print(f"全部完成：{len(result['generated_files'])} 个文件")
    print(f"输出目录：{result['output_dir']}")
    print("=" * 60)


if __name__ == "__main__":
    main()
